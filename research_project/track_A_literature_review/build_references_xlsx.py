"""
build_references_xlsx.py  (Track A/E consolidation)
==================================================

Combines the systematic-review CSV (Track A) and the new-papers CSV (Track E)
into a single, formatted Excel workbook of references / related articles:

  Sheet 1 "Systematic Review"  -- Track A (20 papers, 11 columns)
  Sheet 2 "New Research 2025-26"-- Track E (19 papers, 8 columns)
  Sheet 3 "All References"      -- unified master list (harmonized columns)

Formatting: bold header row, frozen header, auto column widths, wrap text.
"""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RESEARCH_ROOT = Path(__file__).resolve().parents[1]
TRACK_A_CSV = RESEARCH_ROOT / "track_A_literature_review" / "systematic_review.csv"
TRACK_E_CSV = RESEARCH_ROOT / "track_E_new_research" / "new_papers.csv"
OUT_XLSX = RESEARCH_ROOT / "track_A_literature_review" / "related_articles.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def read_csv(path: Path):
    with open(path, "r", encoding="utf-8", newline="") as f:
        rows = list(csv.reader(f))
    return rows[0], rows[1:]


def style_sheet(ws, header, data, wrap_cols=None):
    wrap_cols = wrap_cols or set()
    ws.append(header)
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in data:
        ws.append(row)
    ws.freeze_panes = "A2"
    # Column widths + wrap
    for c in range(1, len(header) + 1):
        col = get_column_letter(c)
        maxlen = max([len(str(header[c - 1]))] +
                     [len(str(r[c - 1])) if c - 1 < len(r) else 0 for r in data])
        ws.column_dimensions[col].width = min(max(12, maxlen // 2 + 4), 60)
        if header[c - 1] in wrap_cols:
            for r in range(2, len(data) + 2):
                ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")


def main():
    wb = Workbook()

    # Sheet 1: Track A
    a_head, a_data = read_csv(TRACK_A_CSV)
    ws1 = wb.active
    ws1.title = "Systematic Review"
    style_sheet(ws1, a_head, a_data,
                wrap_cols={"Brief_Summary", "Problem_Solved", "Methodology_Short"})

    # Sheet 2: Track E
    e_head, e_data = read_csv(TRACK_E_CSV)
    ws2 = wb.create_sheet("New Research 2025-26")
    style_sheet(ws2, e_head, e_data,
                wrap_cols={"Novelty", "Advance_Over_Prior", "Relevance_To_Execution_Phase"})

    # Sheet 3: unified master (harmonized subset)
    master_head = ["Source", "Paper_Title", "Authors", "Year", "Venue", "Link",
                   "Type/Novelty", "Summary/Advance"]
    ws3 = wb.create_sheet("All References")
    master_rows = []
    a_idx = {h: i for i, h in enumerate(a_head)}
    for r in a_data:
        master_rows.append([
            "Track A (SLR)", r[a_idx["Paper_Title"]], r[a_idx["Authors"]],
            r[a_idx["Year"]], r[a_idx["Venue"]], r[a_idx["Published_Link"]],
            r[a_idx["Type"]], r[a_idx["Brief_Summary"]],
        ])
    e_idx = {h: i for i, h in enumerate(e_head)}
    for r in e_data:
        master_rows.append([
            "Track E (2025-26)", r[e_idx["Paper_Title"]], r[e_idx["Authors"]],
            r[e_idx["Year"]], r[e_idx["Venue"]], r[e_idx["Link"]],
            r[e_idx["Novelty"]], r[e_idx["Advance_Over_Prior"]],
        ])
    style_sheet(ws3, master_head, master_rows,
                wrap_cols={"Type/Novelty", "Summary/Advance"})

    wb.save(OUT_XLSX)
    print(f"Wrote {OUT_XLSX}")
    print(f"  Sheet 'Systematic Review': {len(a_data)} papers")
    print(f"  Sheet 'New Research 2025-26': {len(e_data)} papers")
    print(f"  Sheet 'All References': {len(master_rows)} papers total")


if __name__ == "__main__":
    main()
